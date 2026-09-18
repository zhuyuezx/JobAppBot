"""Excel output for a list of stored job rows."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

COLUMNS: list[tuple[str, str, int]] = [
    # header, key into flattened row, column width
    ("First seen", "first_seen", 20),
    ("Posted", "published_at", 20),
    ("Title", "title", 45),
    ("Company", "company", 28),
    ("Location", "location", 32),
    ("Workplace", "workplace_type", 11),
    ("Min YoE", "min_yoe", 9),
    ("Visa flag", "visa", 10),
    ("Clearance", "security_clearance", 14),
    ("Salary min", "yearly_min_comp", 12),
    ("Salary max", "yearly_max_comp", 12),
    ("Category", "category", 22),
    ("Tools", "tools", 40),
    ("Requirements", "requirements_summary", 80),
    ("Apply URL", "apply_url", 60),
    ("hiring.cafe URL", "hc_url", 60),
    ("ATS", "source", 12),
    ("Via", "via", 12),
]


def flatten(row: dict[str, Any]) -> dict[str, Any]:
    job = row["job"]
    return {
        "first_seen": (row.get("first_seen") or "")[:16].replace("T", " "),
        "published_at": (job.get("published_at") or "")[:16].replace("T", " "),
        "title": job.get("title"),
        "company": job.get("company"),
        "location": job.get("location"),
        "workplace_type": job.get("workplace_type"),
        "min_yoe": job.get("min_yoe"),
        "visa": "yes" if job.get("visa_sponsorship") else "",
        "security_clearance": job.get("security_clearance"),
        "yearly_min_comp": job.get("yearly_min_comp"),
        "yearly_max_comp": job.get("yearly_max_comp"),
        "category": job.get("category"),
        "tools": ", ".join(job.get("technical_tools") or []),
        "requirements_summary": job.get("requirements_summary"),
        "apply_url": job.get("apply_url"),
        "hc_url": row.get("hc_url") or job.get("hc_url"),
        "source": job.get("source"),
        "via": job.get("via") or row.get("via") or "hiringcafe",
    }


VIA_LABEL = {"hiringcafe": "hiring.cafe", "simplify": "Simplify", "startupjobs": "startup.jobs"}


def _fill_sheet(ws, rows: list[dict[str, Any]]) -> None:
    ws.append([h for h, _, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        flat = flatten(r)
        ws.append([flat.get(k) for _, k, _ in COLUMNS])
    for i, (_, key, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if key in ("apply_url", "hc_url"):
            for cell in ws[get_column_letter(i)][1:]:
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _stable_group(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Within equal (found day, posted day) buckets, order by company, title, location."""
    out: list[dict[str, Any]] = []
    bucket: list[dict[str, Any]] = []
    key = lambda r: ((r.get("first_seen_date") or (r.get("first_seen") or "")[:10]), (r["job"].get("published_at") or "")[:10])
    for r in rows:
        if bucket and key(bucket[-1]) != key(r):
            out += sorted(bucket, key=lambda x: ((x["job"].get("company") or "").lower(), x["job"].get("title") or "", x["job"].get("location") or ""))
            bucket = []
        bucket.append(r)
    out += sorted(bucket, key=lambda x: ((x["job"].get("company") or "").lower(), x["job"].get("title") or "", x["job"].get("location") or ""))
    return out


def write_excel(rows: list[dict[str, Any]], path: str | Path, sheet_title: str = "jobs") -> Path:
    """One workbook: sheet 'all' (newest discovery first) plus one sheet per source."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # same order as the web page: day found desc, posting day desc, then company / title / location
    rows = sorted(rows, key=lambda r: (r.get("first_seen_date") or (r.get("first_seen") or "")[:10], (r["job"].get("published_at") or "")[:10]), reverse=True)
    rows = _stable_group(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "all"
    _fill_sheet(ws, rows)
    by_via: dict[str, list[dict[str, Any]]] = {}
    for r in rows:
        by_via.setdefault(r["job"].get("via") or r.get("via") or "hiringcafe", []).append(r)
    for via in sorted(by_via, key=lambda v: list(VIA_LABEL).index(v) if v in VIA_LABEL else 99):
        _fill_sheet(wb.create_sheet(VIA_LABEL.get(via, via)[:31]), by_via[via])
    wb.save(path)
    return path
